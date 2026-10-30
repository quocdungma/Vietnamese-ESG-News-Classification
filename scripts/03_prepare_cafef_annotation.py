# Extracted from original Colab notebook. Review paths and execute cells in order.


# === Notebook cell 1 ===
# COLAB COMMAND: !pip install -q \
    pandas \
    numpy \
    openpyxl \
    pyarrow \
    unidecode \
    tqdm

# === Notebook cell 2 ===
from pathlib import Path
from urllib.parse import urlparse

import hashlib
import random
import re
import zipfile

import numpy as np
import pandas as pd

from tqdm.auto import tqdm
from unidecode import unidecode

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Font,
    PatternFill,
    Border,
    Side,
)
from openpyxl.worksheet.datavalidation import (
    DataValidation,
)
from openpyxl.utils import get_column_letter

pd.set_option("display.max_columns", 100)
pd.set_option("display.max_colwidth", 150)

print("Đã import đầy đủ thư viện.")

# === Notebook cell 3 ===
from google.colab import drive

drive.mount("/content/gdrive")

DRIVE_ROOT = Path(
    "/content/gdrive/MyDrive"
)

# ============================================================
# THƯ MỤC CHỨA DỮ LIỆU CRAWL CAFEF
# ============================================================

INPUT_DIR_CANDIDATES = [
    (
        DRIVE_ROOT
        / "Project_Fin_Data_Crawling"
        / "Crawling_news"
        / "Data"
        / "incremental_updates"
    ),
    (
        DRIVE_ROOT
        / "Crawling_news"
        / "Data"
        / "incremental_updates"
    ),
]

CAFEF_INPUT_DIR = next(
    (
        directory
        for directory in INPUT_DIR_CANDIDATES
        if directory.exists()
    ),
    None,
)

if CAFEF_INPUT_DIR is None:
    raise FileNotFoundError(
        "Không tìm thấy thư mục chứa dữ liệu CafeF.\n"
        "Các đường dẫn đã kiểm tra:\n"
        + "\n".join(
            str(path)
            for path in INPUT_DIR_CANDIDATES
        )
    )


# ============================================================
# THƯ MỤC ĐẦU RA RIÊNG CHO CAFEF
# ============================================================

CAFEF_PROJECT_DIR = (
    DRIVE_ROOT
    / "Crawling_news"
    / "ESG_annotation_CafeF_800"
)

SAMPLE_DIR = (
    CAFEF_PROJECT_DIR
    / "01_sample_800"
)

BATCH_DIR = (
    CAFEF_PROJECT_DIR
    / "02_annotation_batches_41_56"
)

REPORT_DIR = (
    CAFEF_PROJECT_DIR
    / "03_reports"
)

for directory in [
    CAFEF_PROJECT_DIR,
    SAMPLE_DIR,
    BATCH_DIR,
    REPORT_DIR,
]:
    directory.mkdir(
        parents=True,
        exist_ok=True,
    )


# ============================================================
# THÔNG SỐ LẤY MẪU
# ============================================================

RANDOM_STATE = 20260814

N_CAFEF_ARTICLES = 800
N_BATCHES = 16
BATCH_SIZE = 50

FIRST_BATCH_NUMBER = 41
LAST_BATCH_NUMBER = (
    FIRST_BATCH_NUMBER
    + N_BATCHES
    - 1
)

MIN_CONTENT_LENGTH = 300
MAX_CONTENT_LENGTH_EXCEL = 30000

assert (
    N_CAFEF_ARTICLES
    == N_BATCHES * BATCH_SIZE
)

print("CafeF input :", CAFEF_INPUT_DIR)
print("Output      :", CAFEF_PROJECT_DIR)
print("Số bài      :", N_CAFEF_ARTICLES)
print("Số batch    :", N_BATCHES)
print(
    "Batch number:",
    FIRST_BATCH_NUMBER,
    "→",
    LAST_BATCH_NUMBER,
)

# === Notebook cell 4 ===
from google.colab import drive

drive.mount("/content/gdrive")

DRIVE_ROOT = Path(
    "/content/gdrive/MyDrive"
)

# ============================================================
# THƯ MỤC CHỨA DỮ LIỆU CRAWL CAFEF
# ============================================================

INPUT_DIR_CANDIDATES = [
    (
        DRIVE_ROOT
        / "Project_Fin_Data_Crawling"
        / "Crawling_news"
        / "Data"
        / "incremental_updates"
    ),
    (
        DRIVE_ROOT
        / "Crawling_news"
        / "Data"
        / "incremental_updates"
    ),
]

CAFEF_INPUT_DIR = next(
    (
        directory
        for directory in INPUT_DIR_CANDIDATES
        if directory.exists()
    ),
    None,
)

if CAFEF_INPUT_DIR is None:
    raise FileNotFoundError(
        "Không tìm thấy thư mục chứa dữ liệu CafeF.\n"
        "Các đường dẫn đã kiểm tra:\n"
        + "\n".join(
            str(path)
            for path in INPUT_DIR_CANDIDATES
        )
    )


# ============================================================
# THƯ MỤC ĐẦU RA RIÊNG CHO CAFEF
# ============================================================

CAFEF_PROJECT_DIR = (
    DRIVE_ROOT
    / "Crawling_news"
    / "ESG_annotation_CafeF_800"
)

SAMPLE_DIR = (
    CAFEF_PROJECT_DIR
    / "01_sample_800"
)

BATCH_DIR = (
    CAFEF_PROJECT_DIR
    / "02_annotation_batches_41_56"
)

REPORT_DIR = (
    CAFEF_PROJECT_DIR
    / "03_reports"
)

for directory in [
    CAFEF_PROJECT_DIR,
    SAMPLE_DIR,
    BATCH_DIR,
    REPORT_DIR,
]:
    directory.mkdir(
        parents=True,
        exist_ok=True,
    )


# ============================================================
# THÔNG SỐ LẤY MẪU
# ============================================================

RANDOM_STATE = 20260814

N_CAFEF_ARTICLES = 800
N_BATCHES = 16
BATCH_SIZE = 50

FIRST_BATCH_NUMBER = 41
LAST_BATCH_NUMBER = (
    FIRST_BATCH_NUMBER
    + N_BATCHES
    - 1
)

MIN_CONTENT_LENGTH = 300
MAX_CONTENT_LENGTH_EXCEL = 30000

assert (
    N_CAFEF_ARTICLES
    == N_BATCHES * BATCH_SIZE
)

print("CafeF input :", CAFEF_INPUT_DIR)
print("Output      :", CAFEF_PROJECT_DIR)
print("Số bài      :", N_CAFEF_ARTICLES)
print("Số batch    :", N_BATCHES)
print(
    "Batch number:",
    FIRST_BATCH_NUMBER,
    "→",
    LAST_BATCH_NUMBER,
)

# === Notebook cell 5 ===
cafef_detail_files = sorted(
    CAFEF_INPUT_DIR.glob(
        "cafef_detail_increment_*.parquet"
    )
)

if not cafef_detail_files:
    raise FileNotFoundError(
        "Không tìm thấy file "
        "cafef_detail_increment_*.parquet trong:\n"
        f"{CAFEF_INPUT_DIR}"
    )

print(
    "Số file CafeF tìm thấy:",
    len(cafef_detail_files),
)

cafef_parts = []

for file_path in cafef_detail_files:
    temp = pd.read_parquet(
        file_path
    )

    temp["input_file"] = (
        file_path.name
    )

    cafef_parts.append(temp)

    print(
        f"{file_path.name:<55}"
        f"{len(temp):>10,} dòng"
    )

cafef_raw = pd.concat(
    cafef_parts,
    ignore_index=True,
)

print(
    "\nTổng số dòng CafeF:",
    f"{len(cafef_raw):,}",
)

print("\nCác cột hiện có:")

display(
    pd.DataFrame(
        {
            "column": (
                cafef_raw.columns.tolist()
            )
        }
    )
)

if "parse_status" in cafef_raw.columns:
    print("\nTrạng thái parser:")

    display(
        cafef_raw["parse_status"]
        .fillna("missing")
        .astype(str)
        .value_counts(dropna=False)
        .to_frame("n_articles")
    )

# === Notebook cell 6 ===
def coalesce_text(
    data,
    candidate_columns,
):
    """
    Lấy nội dung không rỗng đầu tiên
    trong danh sách các cột.
    """
    result = pd.Series(
        "",
        index=data.index,
        dtype="object",
    )

    for column in candidate_columns:
        if column not in data.columns:
            continue

        values = (
            data[column]
            .fillna("")
            .astype(str)
            .str.replace(
                r"\s+",
                " ",
                regex=True,
            )
            .str.strip()
        )

        result = result.mask(
            result.eq(""),
            values,
        )

    return result


def normalize_text_for_duplicate(
    value,
):
    value = unidecode(
        str(value)
    ).lower()

    value = re.sub(
        r"[^a-z0-9]+",
        " ",
        value,
    )

    value = re.sub(
        r"\s+",
        " ",
        value,
    )

    return value.strip()


def make_article_id(
    url,
):
    return hashlib.sha1(
        str(url).encode("utf-8")
    ).hexdigest()[:16]


def convert_to_naive_datetime(
    value,
):
    """
    Chuyển ngày về Timestamp không timezone
    để có thể xuất sang Excel.
    """
    try:
        timestamp = pd.to_datetime(
            value,
            errors="coerce",
        )

        if pd.isna(timestamp):
            return pd.NaT

        timestamp = pd.Timestamp(
            timestamp
        )

        if timestamp.tzinfo is not None:
            timestamp = (
                timestamp
                .tz_convert(
                    "Asia/Ho_Chi_Minh"
                )
                .tz_localize(None)
            )

        return timestamp

    except Exception:
        return pd.NaT


def detect_source_from_url(
    url,
):
    url = str(url).lower()

    if "cafef.vn" in url:
        return "CafeF"

    if "vietstock.vn" in url:
        return "Vietstock"

    return "Unknown"


print("Đã tạo các hàm chuẩn hóa.")

# === Notebook cell 7 ===
cafef_pool = cafef_raw.copy()

# Chỉ giữ những bài được parser đánh dấu thành công
if "parse_status" in cafef_pool.columns:
    status_normalized = (
        cafef_pool["parse_status"]
        .fillna("")
        .astype(str)
        .str.lower()
        .str.strip()
    )

    cafef_pool = cafef_pool[
        status_normalized.eq("success")
    ].copy()


cafef_pool["article_url"] = (
    coalesce_text(
        cafef_pool,
        [
            "url_normalized",
            "url",
            "article_url",
            "link",
        ],
    )
)

cafef_pool["title"] = (
    coalesce_text(
        cafef_pool,
        [
            "title",
            "title_listing",
            "headline",
        ],
    )
)

cafef_pool["description"] = (
    coalesce_text(
        cafef_pool,
        [
            "description",
            "description_listing",
            "sapo",
        ],
    )
)

cafef_pool["content"] = (
    coalesce_text(
        cafef_pool,
        [
            "content",
            "article_content",
            "body",
        ],
    )
)

cafef_pool["category"] = (
    coalesce_text(
        cafef_pool,
        [
            "category",
            "section",
        ],
    )
)

cafef_pool["category"] = (
    cafef_pool["category"]
    .replace(
        "",
        "unknown_category",
    )
)

if "published_at" in cafef_pool.columns:
    cafef_pool["published_at"] = (
        cafef_pool["published_at"]
        .map(
            convert_to_naive_datetime
        )
    )
else:
    cafef_pool["published_at"] = (
        pd.NaT
    )

cafef_pool["source"] = "CafeF"

cafef_pool["article_id"] = (
    cafef_pool["article_url"]
    .map(make_article_id)
)

cafef_pool["combined_text"] = (
    cafef_pool["title"]
    + "\n"
    + cafef_pool["description"]
    + "\n"
    + cafef_pool["content"]
)

cafef_pool["article_length"] = (
    cafef_pool["combined_text"]
    .str.len()
)

cafef_pool["source_from_url"] = (
    cafef_pool["article_url"]
    .map(detect_source_from_url)
)

print(
    "Sau khi giữ parse_status=success:",
    f"{len(cafef_pool):,}",
)

display(
    cafef_pool[
        [
            "source",
            "published_at",
            "category",
            "title",
            "article_length",
            "article_url",
        ]
    ].head()
)

# === Notebook cell 8 ===
n_before_quality = len(
    cafef_pool
)

MAX_ALLOWED_DATE = (
    pd.Timestamp.today()
    .normalize()
    + pd.Timedelta(days=1)
)

valid_date = (
    cafef_pool["published_at"]
    .isna()
    | cafef_pool["published_at"]
    .le(MAX_ALLOWED_DATE)
)

cafef_pool = cafef_pool[
    cafef_pool["source_from_url"]
    .eq("CafeF")
    & cafef_pool["article_url"]
    .str.len()
    .gt(0)
    & cafef_pool["title"]
    .str.len()
    .gt(0)
    & cafef_pool["article_length"]
    .ge(MIN_CONTENT_LENGTH)
    & valid_date
].copy()

print(
    "Trước lọc chất lượng:",
    f"{n_before_quality:,}",
)

print(
    "Sau lọc chất lượng  :",
    f"{len(cafef_pool):,}",
)

print(
    "Đã loại             :",
    f"{n_before_quality - len(cafef_pool):,}",
)

# === Notebook cell 9 ===
# Luôn khởi tạo biến để tránh NameError
vietstock_master = pd.DataFrame()
VIETSTOCK_MASTER_PATH = None

exact_candidates = [
    (
        DRIVE_ROOT
        / "Crawling_news"
        / "ESG_annotation_sampling"
        / "02_sample_2000"
        / "annotation_master_2000.parquet"
    ),
    (
        DRIVE_ROOT
        / "Crawling_news"
        / "ESG_annotation_sampling"
        / "02_sample_2000"
        / "annotation_master_2000.xlsx"
    ),
    (
        DRIVE_ROOT
        / "Project_Fin_Data_Crawling"
        / "Crawling_news"
        / "ESG_annotation_sampling"
        / "02_sample_2000"
        / "annotation_master_2000.parquet"
    ),
    (
        DRIVE_ROOT
        / "Project_Fin_Data_Crawling"
        / "Crawling_news"
        / "ESG_annotation_sampling"
        / "02_sample_2000"
        / "annotation_master_2000.xlsx"
    ),
]

# Ưu tiên các đường dẫn chính xác
VIETSTOCK_MASTER_PATH = next(
    (
        path
        for path in exact_candidates
        if path.exists()
    ),
    None,
)

# Nếu chưa tìm thấy, tìm trong hai thư mục dự án có khả năng nhất
if VIETSTOCK_MASTER_PATH is None:
    search_roots = [
        DRIVE_ROOT / "Crawling_news",
        (
            DRIVE_ROOT
            / "Project_Fin_Data_Crawling"
            / "Crawling_news"
        ),
    ]

    discovered_files = []

    for search_root in search_roots:
        if not search_root.exists():
            continue

        discovered_files.extend(
            search_root.rglob(
                "annotation_master_2000.parquet"
            )
        )

        discovered_files.extend(
            search_root.rglob(
                "annotation_master_2000.xlsx"
            )
        )

    # Loại đường dẫn trùng, ưu tiên parquet
    discovered_files = sorted(
        set(discovered_files),
        key=lambda path: (
            path.suffix.lower() != ".parquet",
            len(path.parts),
            str(path),
        ),
    )

    if discovered_files:
        print("Các master tìm thấy:")

        for path in discovered_files:
            print("-", path)

        VIETSTOCK_MASTER_PATH = (
            discovered_files[0]
        )


# Đọc master nếu tìm thấy
if VIETSTOCK_MASTER_PATH is not None:
    if (
        VIETSTOCK_MASTER_PATH
        .suffix.lower()
        == ".parquet"
    ):
        vietstock_master = pd.read_parquet(
            VIETSTOCK_MASTER_PATH
        )
    else:
        vietstock_master = pd.read_excel(
            VIETSTOCK_MASTER_PATH
        )

    print("\nĐã đọc master Vietstock:")
    print(VIETSTOCK_MASTER_PATH)

    print(
        "Số dòng:",
        f"{len(vietstock_master):,}",
    )

    print(
        "Các cột:",
        vietstock_master.columns.tolist(),
    )

else:
    print(
        "CẢNH BÁO: Không tìm thấy "
        "annotation_master_2000."
    )

    print(
        "Notebook vẫn chạy, nhưng sẽ bỏ qua "
        "kiểm tra trùng nội dung với Vietstock."
    )

print(
    "\nBiến vietstock_master đã được tạo:",
    type(vietstock_master),
)

# === Notebook cell 10 ===
cafef_pool["title_normalized"] = (
    cafef_pool["title"]
    .map(
        normalize_text_for_duplicate
    )
)

cafef_pool[
    "content_prefix_normalized"
] = (
    cafef_pool["content"]
    .str[:1000]
    .map(
        normalize_text_for_duplicate
    )
)

before_internal_dedup = len(
    cafef_pool
)

cafef_pool = (
    cafef_pool
    .sort_values(
        "article_length",
        ascending=False,
    )
    .drop_duplicates(
        subset=["article_url"],
        keep="first",
    )
    .drop_duplicates(
        subset=["title_normalized"],
        keep="first",
    )
    .drop_duplicates(
        subset=[
            "content_prefix_normalized"
        ],
        keep="first",
    )
    .reset_index(drop=True)
)

after_internal_dedup = len(
    cafef_pool
)

print(
    "Trước loại trùng nội bộ:",
    f"{before_internal_dedup:,}",
)

print(
    "Sau loại trùng nội bộ :",
    f"{after_internal_dedup:,}",
)


# ============================================================
# LOẠI TRÙNG VỚI 2.000 BÀI VIETSTOCK
# ============================================================

n_cross_source_duplicates = 0

if not vietstock_master.empty:
    existing_ids = set()

    if (
        "article_id"
        in vietstock_master.columns
    ):
        existing_ids = set(
            vietstock_master[
                "article_id"
            ]
            .dropna()
            .astype(str)
        )

    existing_urls = set()

    for url_column in [
        "article_url",
        "url_normalized",
        "url",
    ]:
        if (
            url_column
            in vietstock_master.columns
        ):
            existing_urls.update(
                vietstock_master[
                    url_column
                ]
                .fillna("")
                .astype(str)
                .str.strip()
            )

    if (
        "title"
        in vietstock_master.columns
    ):
        existing_titles = set(
            vietstock_master["title"]
            .fillna("")
            .map(
                normalize_text_for_duplicate
            )
        )
    else:
        existing_titles = set()

    if (
        "content"
        in vietstock_master.columns
    ):
        existing_content_prefixes = set(
            vietstock_master["content"]
            .fillna("")
            .astype(str)
            .str[:1000]
            .map(
                normalize_text_for_duplicate
            )
        )
    else:
        existing_content_prefixes = set()

    duplicate_with_vietstock = (
        cafef_pool["article_id"]
        .isin(existing_ids)
        | cafef_pool["article_url"]
        .isin(existing_urls)
        | cafef_pool[
            "title_normalized"
        ]
        .isin(existing_titles)
        | cafef_pool[
            "content_prefix_normalized"
        ]
        .isin(
            existing_content_prefixes
        )
    )

    n_cross_source_duplicates = int(
        duplicate_with_vietstock.sum()
    )

    cafef_pool = (
        cafef_pool[
            ~duplicate_with_vietstock
        ]
        .copy()
        .reset_index(drop=True)
    )

print(
    "Số bài trùng Vietstock:",
    f"{n_cross_source_duplicates:,}",
)

print(
    "Pool CafeF cuối cùng :",
    f"{len(cafef_pool):,}",
)

if len(cafef_pool) < N_CAFEF_ARTICLES:
    raise ValueError(
        f"Pool chỉ còn {len(cafef_pool):,} bài, "
        f"không đủ lấy {N_CAFEF_ARTICLES:,} bài."
    )

# === Notebook cell 11 ===
ESG_TITLE_KEYWORDS = [
    # Tổng quát
    "esg",
    "phát triển bền vững",
    "bền vững",
    "trách nhiệm xã hội",

    # Environmental
    "môi trường",
    "ô nhiễm",
    "phát thải",
    "carbon",
    "khí nhà kính",
    "biến đổi khí hậu",
    "năng lượng sạch",
    "năng lượng tái tạo",
    "điện mặt trời",
    "điện gió",
    "chuyển đổi xanh",
    "tăng trưởng xanh",
    "tín chỉ carbon",
    "chất thải",
    "nước thải",
    "tái chế",
    "kinh tế tuần hoàn",
    "đa dạng sinh học",
    "phá rừng",

    # Social
    "người lao động",
    "công nhân",
    "tiền lương",
    "lương tối thiểu",
    "an toàn lao động",
    "tai nạn lao động",
    "phúc lợi",
    "nhân quyền",
    "bình đẳng giới",
    "cộng đồng",
    "từ thiện",
    "an toàn sản phẩm",
    "quyền lợi khách hàng",
    "bảo mật dữ liệu",
    "dữ liệu cá nhân",

    # Governance
    "quản trị doanh nghiệp",
    "hội đồng quản trị",
    "cổ đông",
    "công bố thông tin",
    "kiểm toán",
    "gian lận",
    "tham nhũng",
    "hối lộ",
    "sai phạm",
    "vi phạm",
    "khởi tố",
    "điều tra",
    "xung đột lợi ích",
    "giao dịch bên liên quan",
    "kiểm soát nội bộ",
    "tuân thủ",
]

ESG_FULLTEXT_KEYWORDS = (
    ESG_TITLE_KEYWORDS
    + [
        "quyền cổ đông",
        "cổ đông thiểu số",
        "pha loãng",
        "đạo đức kinh doanh",
        "quản trị rủi ro",
        "rủi ro khí hậu",
        "chuỗi cung ứng bền vững",
        "sản xuất xanh",
        "tiêu dùng xanh",
        "net zero",
        "phát thải ròng bằng 0",
    ]
)

PROMOTION_KEYWORDS = [
    "ưu đãi",
    "khuyến mãi",
    "mở bán",
    "ra mắt sản phẩm",
    "đăng ký ngay",
    "đặt mua",
    "hotline",
    "liên hệ ngay",
    "cơ hội đầu tư hấp dẫn",
    "sản phẩm nổi bật",
    "dịch vụ hàng đầu",
    "khẳng định vị thế",
    "thương hiệu uy tín",
    "giải pháp tối ưu",
    "cam kết mang đến",
    "tiên phong",
    "dẫn đầu thị trường",
]


def build_pattern(
    keywords,
):
    return "|".join(
        re.escape(
            word.lower()
        )
        for word in keywords
    )


esg_title_pattern = build_pattern(
    ESG_TITLE_KEYWORDS
)

esg_full_pattern = build_pattern(
    ESG_FULLTEXT_KEYWORDS
)

promotion_pattern = build_pattern(
    PROMOTION_KEYWORDS
)

print("Đã tạo các bộ từ khóa proxy.")

# === Notebook cell 12 ===
cafef_pool["title_esg_proxy"] = (
    cafef_pool["title"]
    .str.lower()
    .str.contains(
        esg_title_pattern,
        regex=True,
        na=False,
    )
    .astype(int)
)

search_text = (
    cafef_pool["combined_text"]
    .fillna("")
    .astype(str)
    .str.lower()
)

cafef_pool["esg_proxy"] = (
    search_text
    .str.contains(
        esg_full_pattern,
        regex=True,
        na=False,
    )
    .astype(int)
)

cafef_pool["promotion_proxy"] = (
    search_text
    .str.contains(
        promotion_pattern,
        regex=True,
        na=False,
    )
    .astype(int)
)

cafef_pool["sampling_group"] = (
    "ESG_"
    + cafef_pool[
        "esg_proxy"
    ].astype(str)
    + "_PR_"
    + cafef_pool[
        "promotion_proxy"
    ].astype(str)
)

cafef_pool["year_group"] = (
    cafef_pool["published_at"]
    .dt.year
    .fillna(-1)
    .astype(int)
    .astype(str)
)

cafef_pool.loc[
    cafef_pool["year_group"].eq("-1"),
    "year_group",
] = "unknown"

cafef_pool["length_group"] = pd.qcut(
    cafef_pool["article_length"]
    .rank(method="first"),
    q=4,
    labels=[
        "short",
        "medium_short",
        "medium_long",
        "long",
    ],
)

cafef_pool["fine_stratum"] = (
    cafef_pool[
        "sampling_group"
    ].astype(str)
    + "||"
    + cafef_pool[
        "category"
    ].astype(str)
    + "||"
    + cafef_pool[
        "year_group"
    ].astype(str)
    + "||"
    + cafef_pool[
        "length_group"
    ].astype(str)
)

print("Phân phối sampling group:")

display(
    cafef_pool["sampling_group"]
    .value_counts()
    .sort_index()
    .to_frame("available_articles")
)

print("Phân phối category × group:")

display(
    pd.crosstab(
        cafef_pool["category"],
        cafef_pool["sampling_group"],
    )
)

# === Notebook cell 13 ===
def diversified_sample(
    data,
    n,
    stratum_column,
    seed,
):
    if n <= 0:
        return data.iloc[0:0].copy()

    if n >= len(data):
        return data.copy()

    rng = np.random.default_rng(
        seed
    )

    working = data.copy()

    working["_random"] = rng.random(
        len(working)
    )

    working = working.sort_values(
        [
            stratum_column,
            "_random",
        ]
    )

    working["_round"] = (
        working
        .groupby(
            stratum_column,
            observed=True,
        )
        .cumcount()
    )

    result = (
        working
        .sort_values(
            [
                "_round",
                "_random",
            ]
        )
        .head(n)
        .copy()
    )

    return result.drop(
        columns=[
            "_random",
            "_round",
        ],
        errors="ignore",
    )


print("Đã tạo hàm diversified_sample.")

# === Notebook cell 14 ===
# Thiết kế tương đương mẫu Vietstock:
# mục tiêu ban đầu 200 bài cho mỗi nhóm.
GROUP_TARGETS = {
    "ESG_1_PR_1": 200,
    "ESG_1_PR_0": 200,
    "ESG_0_PR_1": 200,
    "ESG_0_PR_0": 200,
}

selected_parts = []

for group_number, (
    group_name,
    target_n,
) in enumerate(
    GROUP_TARGETS.items(),
    start=1,
):
    group_data = cafef_pool[
        cafef_pool[
            "sampling_group"
        ].eq(group_name)
    ].copy()

    actual_n = min(
        target_n,
        len(group_data),
    )

    sampled = diversified_sample(
        group_data,
        n=actual_n,
        stratum_column="fine_stratum",
        seed=(
            RANDOM_STATE
            + 100
            + group_number
        ),
    )

    selected_parts.append(
        sampled
    )

    print(
        f"{group_name}: "
        f"lấy {len(sampled):,}/"
        f"{target_n:,}; "
        f"pool={len(group_data):,}"
    )

sample_df = pd.concat(
    selected_parts,
    ignore_index=True,
).drop_duplicates(
    subset=["article_id"]
)

n_missing = (
    N_CAFEF_ARTICLES
    - len(sample_df)
)

if n_missing > 0:
    remaining_pool = cafef_pool[
        ~cafef_pool["article_id"]
        .isin(
            sample_df["article_id"]
        )
    ].copy()

    if len(remaining_pool) < n_missing:
        raise ValueError(
            f"Cần lấy bù {n_missing:,} bài "
            f"nhưng pool chỉ còn "
            f"{len(remaining_pool):,}."
        )

    fill_sample = diversified_sample(
        remaining_pool,
        n=n_missing,
        stratum_column="fine_stratum",
        seed=RANDOM_STATE + 999,
    )

    sample_df = pd.concat(
        [
            sample_df,
            fill_sample,
        ],
        ignore_index=True,
    )

sample_df = (
    sample_df
    .drop_duplicates(
        subset=["article_id"]
    )
    .sample(
        n=N_CAFEF_ARTICLES,
        random_state=RANDOM_STATE,
    )
    .reset_index(drop=True)
)

sample_df["source"] = "CafeF"

print(
    "\nTổng mẫu CafeF:",
    len(sample_df),
)

display(
    sample_df["sampling_group"]
    .value_counts()
    .sort_index()
    .to_frame("n_articles")
)

assert len(sample_df) == 800
assert (
    sample_df["article_id"]
    .nunique()
    == 800
)
assert (
    sample_df["article_url"]
    .nunique()
    == 800
)

# === Notebook cell 15 ===
def allocate_to_balanced_batches(
    data,
    n_batches,
    batch_size,
    seed,
):
    rng = np.random.default_rng(
        seed
    )

    batches = {
        batch_number: []
        for batch_number in range(
            1,
            n_batches + 1,
        )
    }

    group_counts = {
        batch_number: {}
        for batch_number in batches
    }

    working = (
        data.copy()
        .reset_index(drop=True)
    )

    working["_random"] = rng.random(
        len(working)
    )

    group_sizes = (
        working["sampling_group"]
        .value_counts()
        .to_dict()
    )

    group_order = sorted(
        group_sizes,
        key=group_sizes.get,
        reverse=True,
    )

    for group_name in group_order:
        group_data = working[
            working["sampling_group"]
            .eq(group_name)
        ].sort_values(
            [
                "category",
                "year_group",
                "_random",
            ]
        )

        for row_index in group_data.index:
            available_batches = [
                batch_number
                for batch_number, indices
                in batches.items()
                if len(indices) < batch_size
            ]

            chosen_batch = min(
                available_batches,
                key=lambda batch_number: (
                    group_counts[
                        batch_number
                    ].get(
                        group_name,
                        0,
                    ),
                    len(
                        batches[
                            batch_number
                        ]
                    ),
                    rng.random(),
                ),
            )

            batches[
                chosen_batch
            ].append(row_index)

            group_counts[
                chosen_batch
            ][group_name] = (
                group_counts[
                    chosen_batch
                ].get(
                    group_name,
                    0,
                )
                + 1
            )

    assignments = []

    for batch_number, indices in (
        batches.items()
    ):
        if len(indices) != batch_size:
            raise ValueError(
                f"Batch nội bộ "
                f"{batch_number} có "
                f"{len(indices)} bài."
            )

        for row_index in indices:
            assignments.append(
                {
                    "row_index": (
                        row_index
                    ),
                    "batch_number": (
                        batch_number
                    ),
                }
            )

    assignment_df = (
        pd.DataFrame(assignments)
        .set_index("row_index")
    )

    result = working.join(
        assignment_df,
        how="left",
    )

    return (
        result
        .drop(
            columns="_random"
        )
        .sort_values(
            [
                "batch_number",
                "sampling_group",
            ]
        )
        .reset_index(drop=True)
    )


sample_df = allocate_to_balanced_batches(
    sample_df,
    n_batches=N_BATCHES,
    batch_size=BATCH_SIZE,
    seed=RANDOM_STATE,
)

# Chuyển batch 1–16 thành 41–56
sample_df["batch_number"] = (
    sample_df["batch_number"]
    + FIRST_BATCH_NUMBER
    - 1
)

# Index tiếp nối 2.000 bài Vietstock
sample_df.insert(
    0,
    "index",
    range(
        2001,
        2001 + len(sample_df),
    ),
)

batch_distribution = pd.crosstab(
    sample_df["batch_number"],
    sample_df["sampling_group"],
)

batch_distribution["total"] = (
    batch_distribution.sum(axis=1)
)

display(batch_distribution)

assert (
    batch_distribution["total"]
    .eq(BATCH_SIZE)
    .all()
)

assert (
    sample_df["batch_number"]
    .min()
    == 41
)

assert (
    sample_df["batch_number"]
    .max()
    == 56
)

print(
    "PASS: Đã chia thành batch 41–56."
)

# === Notebook cell 16 ===
ANNOTATION_COLUMNS = [
    "promotion_label",
    "target_type",
    "target_entity_text",
    "target_company_id",
    "target_ticker",
    "sentiment",
    "esg_relevance",
    "environmental",
    "social",
    "governance",
    "esg_impact",
    "esg_information_type",
    "evidence_text",
    "annotator_note",
    "annotator_id",
    "annotation_status",
]

for column in ANNOTATION_COLUMNS:
    sample_df[column] = ""

# Người gán nhãn hiện tại
sample_df["annotator_id"] = 1

sample_df["annotation_status"] = (
    "not_started"
)

DROPDOWN_VALUES = {
    "promotion_label": [
        "independent_editorial",
        "company_sourced",
        "promotional",
        "uncertain",
    ],
    "target_type": [
        "company",
        "industry",
        "market",
        "regulator",
        "project",
        "multiple",
        "other",
        "uncertain",
    ],
    # Chỉ còn 3 lớp; neutral đã gộp vào mixed
    "sentiment": [
        "positive",
        "negative",
        "mixed",
    ],
    "esg_relevance": [
        "relevant",
        "not_relevant",
        "uncertain",
    ],
    "environmental": [
        "0",
        "1",
    ],
    "social": [
        "0",
        "1",
    ],
    "governance": [
        "0",
        "1",
    ],
    "esg_impact": [
        "opportunity",
        "risk",
        "mixed",
        "not_assessable",
        "not_applicable",
    ],
    "esg_information_type": [
        "aspiration",
        "commitment",
        "action",
        "outcome",
        "controversy",
        "mixed",
        "not_applicable",
    ],
    "annotation_status": [
        "not_started",
        "completed",
        "needs_review",
        "adjudicated",
    ],
}

OUTPUT_COLUMNS = [
    "index",
    "article_id",
    "source",
    "published_at",
    "category",
    "title",
    "description",
    "content",
    "article_url",
    "promotion_label",
    "target_type",
    "target_entity_text",
    "target_company_id",
    "target_ticker",
    "sentiment",
    "esg_relevance",
    "environmental",
    "social",
    "governance",
    "esg_impact",
    "esg_information_type",
    "evidence_text",
    "annotator_note",
    "annotator_id",
    "annotation_status",
]

AUDIT_COLUMNS = [
    "article_id",
    "sampling_group",
    "esg_proxy",
    "promotion_proxy",
    "title_esg_proxy",
    "article_length",
    "year_group",
    "length_group",
    "input_file",
    "batch_number",
]

print("Đã tạo schema gán nhãn.")

# === Notebook cell 17 ===
def excel_safe_text(
    value,
    max_length=MAX_CONTENT_LENGTH_EXCEL,
):
    if pd.isna(value):
        return ""

    value = str(value)

    if len(value) > max_length:
        return (
            value[:max_length]
            + "\n[CONTENT TRUNCATED]"
        )

    return value


def create_annotation_workbook(
    batch_df,
    output_path,
):
    workbook = Workbook()

    # ========================================================
    # SHEET 1 — INSTRUCTIONS
    # ========================================================

    instruction_sheet = (
        workbook.active
    )

    instruction_sheet.title = (
        "Instructions"
    )

    instructions = [
        [
            "HƯỚNG DẪN GÁN NHÃN "
            "ESG NEWS — CAFEF"
        ],
        [""],
        [
            "1",
            "Đọc title, description "
            "và content.",
        ],
        [
            "2",
            "Xác định một đối tượng "
            "trung tâm của bài.",
        ],
        [
            "3",
            "Gán promotion_label dựa "
            "trên nguồn và cách trình bày.",
        ],
        [
            "4",
            "Gán sentiment đối với "
            "đối tượng trung tâm.",
        ],
        [
            "5",
            "Sentiment chỉ gồm positive, "
            "negative và mixed.",
        ],
        [
            "6",
            "Neutral được gộp vào mixed.",
        ],
        [
            "7",
            "Xác định ESG relevance "
            "độc lập với sentiment.",
        ],
        [
            "8",
            "Nếu not_relevant, đặt "
            "E=S=G=0.",
        ],
        [
            "9",
            "Nếu not_relevant, đặt "
            "esg_impact và "
            "esg_information_type "
            "là not_applicable.",
        ],
        [
            "10",
            "Sao chép 1–3 câu quan trọng "
            "vào evidence_text.",
        ],
        [
            "11",
            "Giải thích trường hợp khó "
            "trong annotator_note.",
        ],
        [
            "12",
            "Chỉ chọn completed khi đã "
            "hoàn thành toàn bộ nhãn.",
        ],
    ]

    for row in instructions:
        instruction_sheet.append(row)

    instruction_sheet[
        "A1"
    ].font = Font(
        bold=True,
        size=16,
        color="FFFFFF",
    )

    instruction_sheet[
        "A1"
    ].fill = PatternFill(
        "solid",
        fgColor="1F4E78",
    )

    instruction_sheet.column_dimensions[
        "A"
    ].width = 12

    instruction_sheet.column_dimensions[
        "B"
    ].width = 110

    # ========================================================
    # SHEET 2 — ANNOTATION
    # ========================================================

    annotation_sheet = (
        workbook.create_sheet(
            "Annotation"
        )
    )

    for column_index, column_name in enumerate(
        OUTPUT_COLUMNS,
        start=1,
    ):
        cell = annotation_sheet.cell(
            row=1,
            column=column_index,
            value=column_name,
        )

        cell.font = Font(
            bold=True,
            color="FFFFFF",
        )

        cell.fill = PatternFill(
            "solid",
            fgColor="2F75B5",
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    for row_index, (_, row) in enumerate(
        batch_df.iterrows(),
        start=2,
    ):
        for column_index, column_name in enumerate(
            OUTPUT_COLUMNS,
            start=1,
        ):
            value = row.get(
                column_name,
                "",
            )

            if (
                column_name
                == "published_at"
                and pd.notna(value)
            ):
                try:
                    value = pd.Timestamp(
                        value
                    ).strftime(
                        "%Y-%m-%d %H:%M:%S"
                    )
                except Exception:
                    value = str(value)

            annotation_sheet.cell(
                row=row_index,
                column=column_index,
                value=excel_safe_text(
                    value
                ),
            )

    annotation_sheet.freeze_panes = (
        "A2"
    )

    annotation_sheet.auto_filter.ref = (
        annotation_sheet.dimensions
    )

    annotation_sheet.row_dimensions[
        1
    ].height = 35

    widths = {
        "index": 10,
        "article_id": 20,
        "source": 12,
        "published_at": 21,
        "category": 18,
        "title": 48,
        "description": 58,
        "content": 100,
        "article_url": 40,
        "promotion_label": 24,
        "target_type": 18,
        "target_entity_text": 30,
        "target_company_id": 20,
        "target_ticker": 15,
        "sentiment": 16,
        "esg_relevance": 18,
        "environmental": 16,
        "social": 12,
        "governance": 16,
        "esg_impact": 20,
        "esg_information_type": 24,
        "evidence_text": 65,
        "annotator_note": 45,
        "annotator_id": 16,
        "annotation_status": 20,
    }

    for column_index, column_name in enumerate(
        OUTPUT_COLUMNS,
        start=1,
    ):
        annotation_sheet.column_dimensions[
            get_column_letter(
                column_index
            )
        ].width = widths.get(
            column_name,
            18,
        )

    thin_border = Border(
        left=Side(
            style="thin",
            color="D9E2F3",
        ),
        right=Side(
            style="thin",
            color="D9E2F3",
        ),
        top=Side(
            style="thin",
            color="D9E2F3",
        ),
        bottom=Side(
            style="thin",
            color="D9E2F3",
        ),
    )

    for row_cells in (
        annotation_sheet.iter_rows(
            min_row=2,
            max_row=(
                annotation_sheet.max_row
            ),
        )
    ):
        for cell in row_cells:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

            cell.border = thin_border

    # ========================================================
    # SHEET 3 — DROPDOWN VALUES
    # ========================================================

    dropdown_sheet = (
        workbook.create_sheet(
            "Dropdown_Values"
        )
    )

    dropdown_ranges = {}

    for column_index, (
        field,
        values,
    ) in enumerate(
        DROPDOWN_VALUES.items(),
        start=1,
    ):
        dropdown_sheet.cell(
            row=1,
            column=column_index,
            value=field,
        )

        for row_index, value in enumerate(
            values,
            start=2,
        ):
            dropdown_sheet.cell(
                row=row_index,
                column=column_index,
                value=value,
            )

        column_letter = (
            get_column_letter(
                column_index
            )
        )

        dropdown_ranges[field] = (
            f"'Dropdown_Values'!"
            f"${column_letter}$2:"
            f"${column_letter}"
            f"${len(values) + 1}"
        )

    output_positions = {
        name: position + 1
        for position, name
        in enumerate(OUTPUT_COLUMNS)
    }

    for field, formula_range in (
        dropdown_ranges.items()
    ):
        target_letter = (
            get_column_letter(
                output_positions[field]
            )
        )

        validation = DataValidation(
            type="list",
            formula1=formula_range,
            allow_blank=True,
        )

        validation.error = (
            "Hãy chọn giá trị "
            "trong danh sách."
        )

        validation.errorTitle = (
            "Giá trị không hợp lệ"
        )

        annotation_sheet.add_data_validation(
            validation
        )

        validation.add(
            f"{target_letter}2:"
            f"{target_letter}"
            f"{len(batch_df) + 1}"
        )

    dropdown_sheet.sheet_state = (
        "hidden"
    )

    # ========================================================
    # SHEET 4 — SAMPLING AUDIT
    # ========================================================

    audit_sheet = (
        workbook.create_sheet(
            "Sampling_Audit"
        )
    )

    for column_index, name in enumerate(
        AUDIT_COLUMNS,
        start=1,
    ):
        cell = audit_sheet.cell(
            row=1,
            column=column_index,
            value=name,
        )

        cell.font = Font(
            bold=True
        )

    for row_index, (_, row) in enumerate(
        batch_df.iterrows(),
        start=2,
    ):
        for column_index, name in enumerate(
            AUDIT_COLUMNS,
            start=1,
        ):
            audit_sheet.cell(
                row=row_index,
                column=column_index,
                value=excel_safe_text(
                    row.get(
                        name,
                        "",
                    )
                ),
            )

    audit_sheet.sheet_state = (
        "hidden"
    )

    workbook.save(
        output_path
    )


print(
    "Đã tạo hàm "
    "create_annotation_workbook."
)

# === Notebook cell 18 ===
MASTER_PARQUET_PATH = (
    SAMPLE_DIR
    / "annotation_master_cafef_800.parquet"
)

MASTER_EXCEL_PATH = (
    SAMPLE_DIR
    / "annotation_master_cafef_800.xlsx"
)

ZIP_PATH = (
    SAMPLE_DIR
    / "annotation_batches_41_56_cafef.zip"
)

batch_paths = [
    (
        BATCH_DIR
        / (
            f"annotation_batch_"
            f"{batch_number:02d}.xlsx"
        )
    )
    for batch_number in range(
        FIRST_BATCH_NUMBER,
        LAST_BATCH_NUMBER + 1,
    )
]

all_output_paths = (
    batch_paths
    + [
        MASTER_PARQUET_PATH,
        MASTER_EXCEL_PATH,
        ZIP_PATH,
    ]
)

existing_outputs = [
    path
    for path in all_output_paths
    if path.exists()
]

OVERWRITE_EXISTING = False

if (
    existing_outputs
    and not OVERWRITE_EXISTING
):
    raise FileExistsError(
        "Đã có các file đầu ra sau:\n"
        + "\n".join(
            str(path)
            for path in existing_outputs
        )
        + "\n\nNếu chắc chắn muốn tạo lại, "
        "đặt OVERWRITE_EXISTING = True."
    )


# ============================================================
# XUẤT MASTER
# ============================================================

master_export_columns = list(
    dict.fromkeys(
        OUTPUT_COLUMNS
        + [
            "batch_number",
            "sampling_group",
            "esg_proxy",
            "promotion_proxy",
            "title_esg_proxy",
            "article_length",
            "year_group",
            "length_group",
            "input_file",
        ]
    )
)

sample_df.to_parquet(
    MASTER_PARQUET_PATH,
    index=False,
)

sample_df[
    master_export_columns
].to_excel(
    MASTER_EXCEL_PATH,
    index=False,
)


# ============================================================
# XUẤT BATCH 41–56
# ============================================================

created_batch_files = []

for batch_number in tqdm(
    range(
        FIRST_BATCH_NUMBER,
        LAST_BATCH_NUMBER + 1,
    ),
    desc="Creating CafeF batches",
):
    batch_df = sample_df[
        sample_df["batch_number"]
        .eq(batch_number)
    ].copy()

    if len(batch_df) != BATCH_SIZE:
        raise ValueError(
            f"Batch {batch_number} có "
            f"{len(batch_df)} bài, "
            f"không phải {BATCH_SIZE}."
        )

    output_path = (
        BATCH_DIR
        / (
            f"annotation_batch_"
            f"{batch_number:02d}.xlsx"
        )
    )

    create_annotation_workbook(
        batch_df=batch_df,
        output_path=output_path,
    )

    created_batch_files.append(
        output_path
    )


# ============================================================
# TẠO FILE ZIP
# ============================================================

with zipfile.ZipFile(
    ZIP_PATH,
    mode="w",
    compression=zipfile.ZIP_DEFLATED,
) as zip_file:
    for file_path in (
        created_batch_files
    ):
        zip_file.write(
            file_path,
            arcname=file_path.name,
        )

print("\nĐã xuất thành công:")
print("Master parquet:", MASTER_PARQUET_PATH)
print("Master Excel  :", MASTER_EXCEL_PATH)
print("Batch folder  :", BATCH_DIR)
print("ZIP           :", ZIP_PATH)

# === Notebook cell 19 ===
batch_audit_parts = []

for file_path in sorted(
    BATCH_DIR.glob(
        "annotation_batch_*.xlsx"
    )
):
    temp = pd.read_excel(
        file_path,
        sheet_name="Annotation",
        dtype=str,
    )

    temp["batch_file"] = (
        file_path.name
    )

    temp["source_from_url"] = (
        temp["article_url"]
        .map(detect_source_from_url)
    )

    batch_audit_parts.append(
        temp
    )

batch_audit = pd.concat(
    batch_audit_parts,
    ignore_index=True,
)

batch_size_check = (
    batch_audit
    .groupby("batch_file")
    .size()
    .rename("n_articles")
    .reset_index()
)

source_check = (
    batch_audit[
        "source_from_url"
    ]
    .value_counts(dropna=False)
    .rename_axis("source")
    .reset_index(name="n_articles")
)

summary = pd.DataFrame(
    {
        "check": [
            "Tổng số batch",
            "Tổng số bài",
            "Article ID duy nhất",
            "URL duy nhất",
            "Batch nhỏ nhất",
            "Batch lớn nhất",
            "Số bài CafeF theo URL",
        ],
        "value": [
            len(batch_size_check),
            len(batch_audit),
            batch_audit[
                "article_id"
            ].nunique(),
            batch_audit[
                "article_url"
            ].nunique(),
            batch_size_check[
                "n_articles"
            ].min(),
            batch_size_check[
                "n_articles"
            ].max(),
            batch_audit[
                "source_from_url"
            ].eq("CafeF").sum(),
        ],
    }
)

display(summary)
display(batch_size_check)
display(source_check)

assert len(batch_size_check) == 16
assert len(batch_audit) == 800
assert (
    batch_audit["article_id"]
    .nunique()
    == 800
)
assert (
    batch_audit["article_url"]
    .nunique()
    == 800
)
assert (
    batch_size_check["n_articles"]
    .eq(50)
    .all()
)
assert (
    batch_audit["source_from_url"]
    .eq("CafeF")
    .all()
)
assert (
    batch_audit["source"]
    .str.lower()
    .eq("cafef")
    .all()
)

print(
    "PASS: 800 bài CafeF, "
    "16 batch × 50 bài, "
    "batch 41–56."
)

# === Notebook cell 20 ===
