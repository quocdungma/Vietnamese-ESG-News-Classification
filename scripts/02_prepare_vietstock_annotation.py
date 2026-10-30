# Extracted from original Colab notebook. Review paths and execute cells in order.


# === Notebook cell 1 ===
# COLAB COMMAND: !pip install -q \
    pandas \
    numpy \
    openpyxl \
    pyarrow \
    requests \
    beautifulsoup4 \
    lxml \
    trafilatura \
    unidecode \
    tqdm

# === Notebook cell 2 ===
from pathlib import Path
from urllib.parse import urlparse
from concurrent.futures import ThreadPoolExecutor, as_completed

import hashlib
import json
import random
import re
import time
import zipfile

import numpy as np
import pandas as pd
import requests
import trafilatura

from bs4 import BeautifulSoup
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry
from tqdm.auto import tqdm
from unidecode import unidecode

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Font,
    PatternFill,
    Border,
    Side,
)
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

pd.set_option("display.max_columns", 100)
pd.set_option("display.max_colwidth", 150)

print("Đã import đầy đủ thư viện.")

# === Notebook cell 3 ===
from google.colab import drive

drive.mount("/content/gdrive")

# Thư mục gốc trên Google Drive
DRIVE_ROOT = Path("/content/gdrive/MyDrive")

# Thư mục đang chứa các file vietstock_links_*.csv
PROJECT_DIR = DRIVE_ROOT / "Crawling_news"
LINKS_DIR = PROJECT_DIR

# Thư mục lưu toàn bộ kết quả của pipeline
WORK_DIR = PROJECT_DIR / "ESG_annotation_sampling"

CRAWLED_DIR = WORK_DIR / "01_crawled_data"
SAMPLE_DIR = WORK_DIR / "02_sample_2000"
BATCH_DIR = WORK_DIR / "03_annotation_batches"
REPORT_DIR = WORK_DIR / "04_reports"
CHECKPOINT_DIR = WORK_DIR / "checkpoints"

OUTPUT_DIRS = [
    WORK_DIR,
    CRAWLED_DIR,
    SAMPLE_DIR,
    BATCH_DIR,
    REPORT_DIR,
    CHECKPOINT_DIR,
]

for directory in OUTPUT_DIRS:
    directory.mkdir(
        parents=True,
        exist_ok=True,
    )

print("Thư mục chứa link :", LINKS_DIR)
print("Thư mục làm việc  :", WORK_DIR)

print("\nCác thư mục đầu ra:")

for directory in OUTPUT_DIRS:
    print("-", directory)

# === Notebook cell 4 ===
# ============================================================
# 2. CẤU HÌNH
# ============================================================

RANDOM_STATE = 20260801

# Mẫu cuối cùng
N_FINAL_ARTICLES = 2000
N_BATCHES = 40
BATCH_SIZE = 50

# Lấy dư để dự phòng link lỗi, bài trùng hoặc bài quá ngắn
N_CRAWL_CANDIDATES = 3000

# Cố gắng đưa bài có dấu hiệu ESG vào candidate pool
TARGET_TITLE_ESG = 1200
TARGET_TITLE_NON_ESG = N_CRAWL_CANDIDATES - TARGET_TITLE_ESG

# Điều kiện chất lượng toàn văn
MIN_CONTENT_LENGTH = 300
MAX_CONTENT_LENGTH_EXCEL = 30000

# Crawl thận trọng để tránh gây tải cho website
MAX_WORKERS = 4
REQUEST_TIMEOUT = 30
DELAY_MIN = 0.40
DELAY_MAX = 0.90

# Lưu checkpoint sau mỗi số bài
CHECKPOINT_EVERY = 100

assert N_FINAL_ARTICLES == N_BATCHES * BATCH_SIZE

print("Số bài cuối cùng :", N_FINAL_ARTICLES)
print("Số batch         :", N_BATCHES)
print("Số bài mỗi batch :", BATCH_SIZE)
print("Candidate crawl  :", N_CRAWL_CANDIDATES)

# === Notebook cell 5 ===
if not PROJECT_DIR.exists():
    raise FileNotFoundError(
        f"Không tìm thấy thư mục:\n{PROJECT_DIR}\n\n"
        "Hãy kiểm tra lại vị trí Crawling_news trong Google Drive."
    )

link_files = sorted(PROJECT_DIR.glob("vietstock_links_*.csv"))

print(f"Tìm thấy {len(link_files)} file CSV:\n")

for file_path in link_files:
    print("-", file_path.name)

if not link_files:
    raise FileNotFoundError(
        "Không tìm thấy file dạng vietstock_links_*.csv "
        f"trong {PROJECT_DIR}"
    )

# === Notebook cell 6 ===
def read_csv_robust(file_path):
    encodings = [
        "utf-8-sig",
        "utf-8",
        "cp1258",
        "latin1",
    ]

    last_error = None

    for encoding in encodings:
        try:
            return pd.read_csv(
                file_path,
                encoding=encoding
            )
        except Exception as error:
            last_error = error

    raise ValueError(
        f"Không đọc được {file_path.name}: {last_error}"
    )

# === Notebook cell 7 ===
link_dfs = []

for file_path in link_files:
    temp = read_csv_robust(file_path)

    temp.columns = [
        unidecode(str(col))
        .lower()
        .strip()
        .replace(" ", "_")
        for col in temp.columns
    ]

    temp["input_file"] = file_path.name
    link_dfs.append(temp)

    print(
        f"{file_path.name:<45} "
        f"{len(temp):>8,} dòng | "
        f"cột: {temp.columns.tolist()}"
    )

links_raw = pd.concat(
    link_dfs,
    ignore_index=True
)

print("\nTổng số dòng:", len(links_raw))
display(links_raw.head())

# === Notebook cell 8 ===
COLUMN_ALIASES = {
    "title": [
        "title",
        "headline",
        "news_title",
        "tieu_de",
    ],
    "article_url": [
        "article_url",
        "url",
        "link",
        "news_url",
    ],
    "category": [
        "category",
        "section",
        "topic",
        "chuyen_muc",
    ],
}


def find_column(columns, aliases):
    for alias in aliases:
        if alias in columns:
            return alias
    return None


rename_mapping = {}

for standard_name, aliases in COLUMN_ALIASES.items():
    found_column = find_column(
        links_raw.columns,
        aliases
    )

    if found_column is not None:
        rename_mapping[found_column] = standard_name


links = links_raw.rename(
    columns=rename_mapping
).copy()

required_columns = [
    "title",
    "article_url",
    "category",
]

missing_columns = [
    col for col in required_columns
    if col not in links.columns
]

if missing_columns:
    raise ValueError(
        f"Thiếu các cột bắt buộc: {missing_columns}\n"
        f"Các cột đang có: {links.columns.tolist()}"
    )

for col in required_columns:
    links[col] = (
        links[col]
        .fillna("")
        .astype(str)
        .str.replace(r"\s+", " ", regex=True)
        .str.strip()
    )

print("Mapping:", rename_mapping)
display(links[required_columns].head())

# === Notebook cell 9 ===
def is_valid_http_url(value):
    try:
        parsed = urlparse(str(value))
        return parsed.scheme in {"http", "https"} and bool(parsed.netloc)
    except Exception:
        return False


links["valid_url"] = links["article_url"].map(
    is_valid_http_url
)

invalid_links = links.loc[
    ~links["valid_url"],
    ["title", "article_url", "input_file"]
].copy()

print("Tổng dòng ban đầu :", len(links))
print("Link không hợp lệ :", len(invalid_links))

links = links[
    links["valid_url"]
    & links["title"].str.len().gt(0)
].copy()

links = (
    links
    .drop_duplicates(subset=["article_url"])
    .reset_index(drop=True)
)

print("Link hợp lệ, không trùng:", len(links))

# === Notebook cell 10 ===
def make_article_id(url):
    return hashlib.sha1(
        str(url).encode("utf-8")
    ).hexdigest()[:16]


def extract_url_year(url):
    match = re.search(
        r"vietstock\.vn/(\d{4})/",
        str(url)
    )
    return int(match.group(1)) if match else np.nan


def extract_url_month(url):
    match = re.search(
        r"vietstock\.vn/\d{4}/(\d{2})/",
        str(url)
    )
    return int(match.group(1)) if match else np.nan


links["article_id"] = links["article_url"].map(
    make_article_id
)

links["url_year"] = links["article_url"].map(
    extract_url_year
)

links["url_month"] = links["article_url"].map(
    extract_url_month
)

links["source"] = "vietstock"

print("Phân phối theo năm:")
display(
    links["url_year"]
    .value_counts(dropna=False)
    .sort_index()
)

# === Notebook cell 11 ===
ESG_TITLE_KEYWORDS = [
    # Tổng quát
    "esg",
    "phát triển bền vững",
    "bền vững",
    "trách nhiệm xã hội",

    # Environmental
    "môi trường",
    "ô nhiễm",
    "phát thải",
    "carbon",
    "khí nhà kính",
    "biến đổi khí hậu",
    "năng lượng sạch",
    "năng lượng tái tạo",
    "điện mặt trời",
    "điện gió",
    "chuyển đổi xanh",
    "tăng trưởng xanh",
    "tín chỉ carbon",
    "chất thải",
    "nước thải",
    "tái chế",
    "kinh tế tuần hoàn",
    "đa dạng sinh học",
    "phá rừng",

    # Social
    "người lao động",
    "công nhân",
    "tiền lương",
    "lương tối thiểu",
    "an toàn lao động",
    "tai nạn lao động",
    "phúc lợi",
    "nhân quyền",
    "bình đẳng giới",
    "cộng đồng",
    "từ thiện",
    "an toàn sản phẩm",
    "quyền lợi khách hàng",
    "bảo mật dữ liệu",
    "dữ liệu cá nhân",

    # Governance
    "quản trị doanh nghiệp",
    "hội đồng quản trị",
    "cổ đông",
    "công bố thông tin",
    "kiểm toán",
    "gian lận",
    "tham nhũng",
    "hối lộ",
    "sai phạm",
    "vi phạm",
    "khởi tố",
    "điều tra",
    "xung đột lợi ích",
    "giao dịch bên liên quan",
    "kiểm soát nội bộ",
    "tuân thủ",
]


def build_pattern(keywords):
    return "|".join(
        re.escape(word.lower())
        for word in keywords
    )


esg_title_pattern = build_pattern(
    ESG_TITLE_KEYWORDS
)

links["title_esg_proxy"] = (
    links["title"]
    .str.lower()
    .str.contains(
        esg_title_pattern,
        regex=True,
        na=False,
    )
    .astype(int)
)

display(
    links["title_esg_proxy"]
    .value_counts()
    .rename_axis("title_esg_proxy")
    .to_frame("n_articles")
)

# === Notebook cell 12 ===
links["year_group"] = (
    links["url_year"]
    .fillna(-1)
    .astype(int)
    .astype(str)
)

links.loc[
    links["year_group"].eq("-1"),
    "year_group"
] = "unknown"

links["category"] = links["category"].replace(
    "",
    "unknown_category"
)

links["candidate_stratum"] = (
    links["title_esg_proxy"].astype(str)
    + "||"
    + links["category"].astype(str)
    + "||"
    + links["year_group"].astype(str)
    + "||"
    + links["input_file"].astype(str)
)

print(
    "Số strata:",
    links["candidate_stratum"].nunique()
)

# === Notebook cell 13 ===
def diversified_sample(
    data,
    n,
    stratum_column,
    seed,
):
    if n <= 0:
        return data.iloc[0:0].copy()

    if n >= len(data):
        return data.copy()

    rng = np.random.default_rng(seed)

    working = data.copy()
    working["_random"] = rng.random(len(working))

    remaining = working.sort_values(
        [stratum_column, "_random"]
    )

    selected_indices = []

    while (
        len(selected_indices) < n
        and len(remaining) > 0
    ):
        round_sample = (
            remaining
            .groupby(
                stratum_column,
                group_keys=False,
                observed=True,
            )
            .head(1)
        )

        needed = n - len(selected_indices)

        if len(round_sample) > needed:
            round_sample = round_sample.sample(
                n=needed,
                random_state=seed + len(selected_indices),
            )

        selected_indices.extend(
            round_sample.index.tolist()
        )

        remaining = remaining.drop(
            round_sample.index
        )

    result = working.loc[selected_indices].copy()

    return result.drop(
        columns="_random",
        errors="ignore"
    )

# === Notebook cell 14 ===
esg_pool = links[
    links["title_esg_proxy"].eq(1)
].copy()

non_esg_pool = links[
    links["title_esg_proxy"].eq(0)
].copy()

n_esg = min(
    TARGET_TITLE_ESG,
    len(esg_pool)
)

candidate_esg = diversified_sample(
    esg_pool,
    n=n_esg,
    stratum_column="candidate_stratum",
    seed=RANDOM_STATE + 1,
)

candidate_non_esg = diversified_sample(
    non_esg_pool,
    n=min(
        TARGET_TITLE_NON_ESG,
        len(non_esg_pool),
    ),
    stratum_column="candidate_stratum",
    seed=RANDOM_STATE + 2,
)

candidates = pd.concat(
    [candidate_esg, candidate_non_esg],
    ignore_index=True,
).drop_duplicates("article_id")

n_missing = N_CRAWL_CANDIDATES - len(candidates)

if n_missing > 0:
    remaining_pool = links[
        ~links["article_id"].isin(
            candidates["article_id"]
        )
    ].copy()

    fill_sample = diversified_sample(
        remaining_pool,
        n=min(n_missing, len(remaining_pool)),
        stratum_column="candidate_stratum",
        seed=RANDOM_STATE + 3,
    )

    candidates = pd.concat(
        [candidates, fill_sample],
        ignore_index=True,
    )

candidates = (
    candidates
    .drop_duplicates("article_id")
    .sample(
        frac=1,
        random_state=RANDOM_STATE,
    )
    .reset_index(drop=True)
)

print("Candidate ESG proxy     :", len(candidate_esg))
print("Candidate non-ESG proxy :", len(candidate_non_esg))
print("Tổng candidate          :", len(candidates))

display(
    pd.crosstab(
        candidates["category"],
        candidates["title_esg_proxy"],
    )
)

# === Notebook cell 15 ===
USER_AGENTS = [
    (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/150.0.0.0 Safari/537.36"
    ),
    (
        "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/149.0.0.0 Safari/537.36"
    ),
]


def create_session():
    session = requests.Session()

    retry_strategy = Retry(
        total=3,
        connect=3,
        read=3,
        backoff_factor=1.5,
        status_forcelist=[
            429,
            500,
            502,
            503,
            504,
        ],
        allowed_methods=["GET"],
    )

    adapter = HTTPAdapter(
        max_retries=retry_strategy
    )

    session.mount("http://", adapter)
    session.mount("https://", adapter)

    return session

# === Notebook cell 16 ===
def get_meta_content(soup, attributes):
    for attrs in attributes:
        tag = soup.find("meta", attrs=attrs)

        if tag and tag.get("content"):
            return tag["content"].strip()

    return ""


def extract_json_ld(soup):
    records = []

    scripts = soup.find_all(
        "script",
        attrs={"type": "application/ld+json"},
    )

    for script in scripts:
        raw_text = script.string or script.get_text()

        if not raw_text:
            continue

        try:
            parsed = json.loads(raw_text)
        except Exception:
            continue

        if isinstance(parsed, list):
            records.extend(parsed)
        elif isinstance(parsed, dict):
            if isinstance(parsed.get("@graph"), list):
                records.extend(parsed["@graph"])
            else:
                records.append(parsed)

    return records


def find_json_ld_field(records, field):
    valid_types = {
        "NewsArticle",
        "Article",
        "ReportageNewsArticle",
    }

    for record in records:
        if not isinstance(record, dict):
            continue

        record_type = record.get("@type")

        if (
            record_type in valid_types
            or (
                isinstance(record_type, list)
                and any(x in valid_types for x in record_type)
            )
        ):
            value = record.get(field)

            if value:
                return value

    return ""

# === Notebook cell 17 ===
def clean_extracted_text(value):
    if value is None:
        return ""

    value = str(value)
    value = re.sub(r"\r\n?", "\n", value)
    value = re.sub(r"[ \t]+", " ", value)
    value = re.sub(r"\n{3,}", "\n\n", value)

    return value.strip()


def extract_content_with_trafilatura(
    html,
    url,
):
    content = trafilatura.extract(
        html,
        url=url,
        include_comments=False,
        include_tables=False,
        include_links=False,
        favor_precision=True,
        deduplicate=True,
        output_format="txt",
    )

    return clean_extracted_text(content)


def extract_content_fallback(soup):
    selectors = [
        "div.article-content",
        "div.article__content",
        "div#vst_detail",
        "div#article-content",
        "div.content-detail",
        "div.detail-content",
        "article",
    ]

    for selector in selectors:
        element = soup.select_one(selector)

        if element is None:
            continue

        for unwanted in element.select(
            "script, style, iframe, "
            ".related-news, .social-share, "
            ".advertisement, .ads"
        ):
            unwanted.decompose()

        paragraphs = [
            p.get_text(" ", strip=True)
            for p in element.find_all(
                ["p", "h2", "h3", "li"]
            )
        ]

        text = "\n".join(
            paragraph
            for paragraph in paragraphs
            if paragraph
        )

        text = clean_extracted_text(text)

        if len(text) >= 200:
            return text

    return ""

# === Notebook cell 18 ===
def crawl_one_article(row):
    article_id = row["article_id"]
    url = row["article_url"]

    result = {
        "article_id": article_id,
        "article_url": url,
        "title_link": row.get("title", ""),
        "category": row.get("category", ""),
        "input_file": row.get("input_file", ""),
        "url_year": row.get("url_year", np.nan),
        "url_month": row.get("url_month", np.nan),
        "title_esg_proxy": row.get(
            "title_esg_proxy",
            0,
        ),
        "crawl_status": "failed",
        "http_status": np.nan,
        "title": "",
        "description": "",
        "published_at": "",
        "author": "",
        "content": "",
        "content_length": 0,
        "crawl_error": "",
    }

    time.sleep(
        random.uniform(
            DELAY_MIN,
            DELAY_MAX,
        )
    )

    try:
        session = create_session()

        headers = {
            "User-Agent": random.choice(USER_AGENTS),
            "Accept": (
                "text/html,application/xhtml+xml,"
                "application/xml;q=0.9,*/*;q=0.8"
            ),
            "Accept-Language": "vi-VN,vi;q=0.9,en;q=0.7",
            "Referer": "https://vietstock.vn/",
            "Connection": "keep-alive",
        }

        response = session.get(
            url,
            headers=headers,
            timeout=REQUEST_TIMEOUT,
        )

        result["http_status"] = response.status_code

        if response.status_code != 200:
            result["crawl_error"] = (
                f"HTTP {response.status_code}"
            )
            return result

        response.encoding = (
            response.apparent_encoding
            or response.encoding
            or "utf-8"
        )

        html = response.text
        soup = BeautifulSoup(html, "lxml")
        json_ld_records = extract_json_ld(soup)

        title = get_meta_content(
            soup,
            [
                {"property": "og:title"},
                {"name": "twitter:title"},
            ],
        )

        if not title:
            title = find_json_ld_field(
                json_ld_records,
                "headline",
            )

        if not title and soup.title:
            title = soup.title.get_text(
                " ",
                strip=True,
            )

        description = get_meta_content(
            soup,
            [
                {"property": "og:description"},
                {"name": "description"},
                {"name": "twitter:description"},
            ],
        )

        if not description:
            description = find_json_ld_field(
                json_ld_records,
                "description",
            )

        published_at = find_json_ld_field(
            json_ld_records,
            "datePublished",
        )

        if not published_at:
            published_at = get_meta_content(
                soup,
                [
                    {
                        "property":
                        "article:published_time"
                    },
                    {"name": "pubdate"},
                    {"name": "date"},
                ],
            )

        author = find_json_ld_field(
            json_ld_records,
            "author",
        )

        if isinstance(author, dict):
            author = author.get("name", "")
        elif isinstance(author, list):
            author = "; ".join(
                item.get("name", "")
                if isinstance(item, dict)
                else str(item)
                for item in author
            )

        content = extract_content_with_trafilatura(
            html,
            url,
        )

        if len(content) < 200:
            content = extract_content_fallback(
                soup
            )

        result.update({
            "title": (
                clean_extracted_text(title)
                or row.get("title", "")
            ),
            "description":
                clean_extracted_text(description),
            "published_at":
                clean_extracted_text(published_at),
            "author":
                clean_extracted_text(author),
            "content": content,
            "content_length": len(content),
        })

        if len(content) >= MIN_CONTENT_LENGTH:
            result["crawl_status"] = "success"
        elif len(content) > 0:
            result["crawl_status"] = "too_short"
            result["crawl_error"] = (
                "Content shorter than minimum"
            )
        else:
            result["crawl_status"] = "no_content"
            result["crawl_error"] = (
                "Could not extract article content"
            )

        return result

    except Exception as error:
        result["crawl_error"] = (
            f"{type(error).__name__}: {error}"
        )
        return result

# === Notebook cell 19 ===
test_rows = candidates.head(3)

test_results = []

for _, row in tqdm(
    test_rows.iterrows(),
    total=len(test_rows),
    desc="Test crawl",
):
    test_results.append(
        crawl_one_article(row)
    )

test_df = pd.DataFrame(test_results)

display(
    test_df[
        [
            "article_id",
            "http_status",
            "crawl_status",
            "title",
            "published_at",
            "content_length",
            "crawl_error",
        ]
    ]
)

for i, row in test_df.iterrows():
    print("=" * 100)
    print("TITLE:", row["title"])
    print("URL:", row["article_url"])
    print("STATUS:", row["crawl_status"])
    print(row["content"][:1000])

# === Notebook cell 20 ===
MASTER_CHECKPOINT = (
    CHECKPOINT_DIR
    / "vietstock_crawl_checkpoint.parquet"
)

if MASTER_CHECKPOINT.exists():
    crawled_existing = pd.read_parquet(
        MASTER_CHECKPOINT
    )

    completed_ids = set(
        crawled_existing["article_id"]
        .dropna()
        .astype(str)
    )

    print(
        "Checkpoint hiện có:",
        len(crawled_existing),
        "bài",
    )
else:
    crawled_existing = pd.DataFrame()
    completed_ids = set()

    print("Chưa có checkpoint cũ.")

remaining_candidates = candidates[
    ~candidates["article_id"].isin(
        completed_ids
    )
].copy()

print("Còn phải crawl:", len(remaining_candidates))

# === Notebook cell 21 ===
new_results = []
all_checkpoint_parts = []

if not crawled_existing.empty:
    all_checkpoint_parts.append(
        crawled_existing
    )

rows_to_crawl = [
    row
    for _, row in remaining_candidates.iterrows()
]

with ThreadPoolExecutor(
    max_workers=MAX_WORKERS
) as executor:

    future_mapping = {
        executor.submit(
            crawl_one_article,
            row,
        ): row["article_id"]
        for row in rows_to_crawl
    }

    progress_bar = tqdm(
        total=len(future_mapping),
        desc="Crawling Vietstock",
    )

    for future in as_completed(future_mapping):
        article_id = future_mapping[future]

        try:
            result = future.result()
        except Exception as error:
            result = {
                "article_id": article_id,
                "crawl_status": "failed",
                "crawl_error": str(error),
            }

        new_results.append(result)
        progress_bar.update(1)

        if (
            len(new_results)
# COLAB COMMAND:             % CHECKPOINT_EVERY
            == 0
        ):
            new_df = pd.DataFrame(
                new_results
            )

            checkpoint_df = pd.concat(
                all_checkpoint_parts
                + [new_df],
                ignore_index=True,
            )

            checkpoint_df = (
                checkpoint_df
                .drop_duplicates(
                    subset=["article_id"],
                    keep="last",
                )
            )

            checkpoint_df.to_parquet(
                MASTER_CHECKPOINT,
                index=False,
            )

            print(
                f"\nCheckpoint: "
                f"{len(checkpoint_df):,} bài"
            )

    progress_bar.close()

# === Notebook cell 22 ===
new_crawled_df = pd.DataFrame(
    new_results
)

crawled_df = pd.concat(
    [
        crawled_existing,
        new_crawled_df,
    ],
    ignore_index=True,
)

crawled_df = (
    crawled_df
    .drop_duplicates(
        subset=["article_id"],
        keep="last",
    )
    .reset_index(drop=True)
)

crawled_df.to_parquet(
    MASTER_CHECKPOINT,
    index=False,
)

full_crawl_parquet = (
    CRAWLED_DIR
    / "vietstock_crawled_candidates.parquet"
)

full_crawl_csv = (
    CRAWLED_DIR
    / "vietstock_crawled_candidates.csv"
)

crawled_df.to_parquet(
    full_crawl_parquet,
    index=False,
)

crawled_df.to_csv(
    full_crawl_csv,
    index=False,
    encoding="utf-8-sig",
)

print("Tổng kết quả:", len(crawled_df))
print("\nTrạng thái crawl:")

display(
    crawled_df["crawl_status"]
    .value_counts(dropna=False)
    .to_frame("n_articles")
)

# === Notebook cell 23 ===
crawl_audit = (
    crawled_df
    .groupby(
        [
            "crawl_status",
            "http_status",
            "crawl_error",
        ],
        dropna=False,
    )
    .size()
    .reset_index(name="n_articles")
    .sort_values(
        "n_articles",
        ascending=False,
    )
)

display(crawl_audit.head(30))

crawl_audit.to_excel(
    REPORT_DIR / "crawl_audit.xlsx",
    index=False,
)

# === Notebook cell 24 ===
articles = crawled_df[
    crawled_df["crawl_status"].eq("success")
].copy()

articles["title"] = (
    articles["title"]
    .fillna("")
    .astype(str)
    .str.replace(r"\s+", " ", regex=True)
    .str.strip()
)

articles["description"] = (
    articles["description"]
    .fillna("")
    .astype(str)
    .str.replace(r"\s+", " ", regex=True)
    .str.strip()
)

articles["content"] = (
    articles["content"]
    .fillna("")
    .astype(str)
    .str.replace(r"[ \t]+", " ", regex=True)
    .str.replace(r"\n{3,}", "\n\n", regex=True)
    .str.strip()
)

articles["published_at"] = pd.to_datetime(
    articles["published_at"],
    errors="coerce",
)

articles["combined_text"] = (
    articles["title"]
    + "\n"
    + articles["description"]
    + "\n"
    + articles["content"]
)

articles["article_length"] = (
    articles["combined_text"].str.len()
)

articles = articles[
    articles["article_length"].ge(
        MIN_CONTENT_LENGTH
    )
].copy()

print("Số bài hợp lệ:", len(articles))

# === Notebook cell 25 ===
def normalize_text_for_duplicate(value):
    value = unidecode(str(value)).lower()
    value = re.sub(r"[^a-z0-9]+", " ", value)
    value = re.sub(r"\s+", " ", value)
    return value.strip()


articles["title_normalized"] = (
    articles["title"].map(
        normalize_text_for_duplicate
    )
)

articles["content_prefix_normalized"] = (
    articles["content"]
    .str[:1000]
    .map(normalize_text_for_duplicate)
)

before_dedup = len(articles)

articles = (
    articles
    .sort_values(
        "article_length",
        ascending=False,
    )
    .drop_duplicates(
        subset=["article_url"]
    )
    .drop_duplicates(
        subset=["title_normalized"]
    )
    .drop_duplicates(
        subset=["content_prefix_normalized"]
    )
    .reset_index(drop=True)
)

print("Trước loại trùng:", before_dedup)
print("Sau loại trùng  :", len(articles))
print("Đã loại         :", before_dedup - len(articles))

# === Notebook cell 27 ===
ESG_FULLTEXT_KEYWORDS = ESG_TITLE_KEYWORDS + [
    "quyền cổ đông",
    "cổ đông thiểu số",
    "pha loãng",
    "đạo đức kinh doanh",
    "quản trị rủi ro",
    "rủi ro khí hậu",
    "chuỗi cung ứng bền vững",
    "sản xuất xanh",
    "tiêu dùng xanh",
    "net zero",
    "phát thải ròng bằng 0",
]

PROMOTION_KEYWORDS = [
    "ưu đãi",
    "khuyến mãi",
    "mở bán",
    "ra mắt sản phẩm",
    "đăng ký ngay",
    "đặt mua",
    "hotline",
    "liên hệ ngay",
    "cơ hội đầu tư hấp dẫn",
    "sản phẩm nổi bật",
    "dịch vụ hàng đầu",
    "khẳng định vị thế",
    "thương hiệu uy tín",
    "giải pháp tối ưu",
    "cam kết mang đến",
    "tiên phong",
    "dẫn đầu thị trường",
]

esg_full_pattern = build_pattern(
    ESG_FULLTEXT_KEYWORDS
)

promotion_pattern = build_pattern(
    PROMOTION_KEYWORDS
)

search_text = articles[
    "combined_text"
].str.lower()

articles["esg_proxy"] = (
    search_text
    .str.contains(
        esg_full_pattern,
        regex=True,
        na=False,
    )
    .astype(int)
)

articles["promotion_proxy"] = (
    search_text
    .str.contains(
        promotion_pattern,
        regex=True,
        na=False,
    )
    .astype(int)
)

articles["sampling_group"] = (
    "ESG_"
    + articles["esg_proxy"].astype(str)
    + "_PR_"
    + articles["promotion_proxy"].astype(str)
)

display(
    articles["sampling_group"]
    .value_counts()
    .sort_index()
    .to_frame("available_articles")
)

# === Notebook cell 28 ===
articles["year_group"] = (
    articles["published_at"]
    .dt.year
    .fillna(articles["url_year"])
    .fillna(-1)
    .astype(int)
    .astype(str)
)

articles.loc[
    articles["year_group"].eq("-1"),
    "year_group"
] = "unknown"

articles["length_group"] = pd.qcut(
    articles["article_length"].rank(
        method="first"
    ),
    q=4,
    labels=[
        "short",
        "medium_short",
        "medium_long",
        "long",
    ],
)

articles["fine_stratum"] = (
    articles["sampling_group"].astype(str)
    + "||"
    + articles["category"].astype(str)
    + "||"
    + articles["year_group"].astype(str)
    + "||"
    + articles["length_group"].astype(str)
)

# === Notebook cell 29 ===
GROUP_TARGETS = {
    "ESG_1_PR_1": 500,
    "ESG_1_PR_0": 500,
    "ESG_0_PR_1": 500,
    "ESG_0_PR_0": 500,
}

selected_parts = []

for group_number, (
    group_name,
    target_n,
) in enumerate(
    GROUP_TARGETS.items(),
    start=1,
):
    group_data = articles[
        articles["sampling_group"].eq(
            group_name
        )
    ].copy()

    actual_n = min(
        target_n,
        len(group_data),
    )

    sampled = diversified_sample(
        group_data,
        n=actual_n,
        stratum_column="fine_stratum",
        seed=RANDOM_STATE + 100 + group_number,
    )

    selected_parts.append(sampled)

    print(
        f"{group_name}: "
        f"lấy {len(sampled):,}/"
        f"{target_n:,}; "
        f"pool={len(group_data):,}"
    )

sample_df = pd.concat(
    selected_parts,
    ignore_index=True,
).drop_duplicates("article_id")

n_missing = (
    N_FINAL_ARTICLES
    - len(sample_df)
)

if n_missing > 0:
    remaining_pool = articles[
        ~articles["article_id"].isin(
            sample_df["article_id"]
        )
    ].copy()

    if len(remaining_pool) < n_missing:
        raise ValueError(
            f"Thiếu dữ liệu: cần thêm {n_missing:,} bài "
            f"nhưng pool chỉ còn {len(remaining_pool):,}."
        )

    fill_sample = diversified_sample(
        remaining_pool,
        n=n_missing,
        stratum_column="fine_stratum",
        seed=RANDOM_STATE + 999,
    )

    sample_df = pd.concat(
        [sample_df, fill_sample],
        ignore_index=True,
    )

sample_df = (
    sample_df
    .drop_duplicates("article_id")
    .sample(
        n=N_FINAL_ARTICLES,
        random_state=RANDOM_STATE,
    )
    .reset_index(drop=True)
)

print("\nTổng mẫu chính thức:", len(sample_df))

display(
    sample_df["sampling_group"]
    .value_counts()
    .to_frame("n_articles")
)

# === Notebook cell 30 ===
def allocate_to_balanced_batches(
    data,
    n_batches,
    batch_size,
    seed,
):
    rng = np.random.default_rng(seed)

    batches = {
        batch_number: []
        for batch_number in range(
            1,
            n_batches + 1
        )
    }

    group_counts = {
        batch_number: {}
        for batch_number in batches
    }

    working = data.copy()
    working["_random"] = rng.random(
        len(working)
    )

    group_sizes = (
        working["sampling_group"]
        .value_counts()
        .to_dict()
    )

    group_order = sorted(
        group_sizes,
        key=group_sizes.get,
        reverse=True,
    )

    for group_name in group_order:
        group_data = working[
            working["sampling_group"].eq(
                group_name
            )
        ].sort_values(
            [
                "category",
                "year_group",
                "_random",
            ]
        )

        for row_index in group_data.index:
            available_batches = [
                batch_number
                for batch_number, indices
                in batches.items()
                if len(indices) < batch_size
            ]

            chosen_batch = min(
                available_batches,
                key=lambda batch_number: (
                    group_counts[
                        batch_number
                    ].get(group_name, 0),
                    len(batches[batch_number]),
                    rng.random(),
                ),
            )

            batches[chosen_batch].append(
                row_index
            )

            group_counts[chosen_batch][
                group_name
            ] = (
                group_counts[
                    chosen_batch
                ].get(group_name, 0)
                + 1
            )

    assignments = []

    for batch_number, indices in batches.items():
        if len(indices) != batch_size:
            raise ValueError(
                f"Batch {batch_number} có "
                f"{len(indices)} bài."
            )

        for row_index in indices:
            assignments.append(
                {
                    "row_index": row_index,
                    "batch_number": batch_number,
                }
            )

    assignment_df = pd.DataFrame(
        assignments
    ).set_index("row_index")

    result = working.join(
        assignment_df,
        how="left",
    )

    return (
        result
        .drop(columns="_random")
        .sort_values(
            [
                "batch_number",
                "sampling_group",
            ]
        )
        .reset_index(drop=True)
    )


sample_df = allocate_to_balanced_batches(
    sample_df,
    n_batches=N_BATCHES,
    batch_size=BATCH_SIZE,
    seed=RANDOM_STATE,
)

display(
    sample_df["batch_number"]
    .value_counts()
    .sort_index()
    .to_frame("n_articles")
)

# === Notebook cell 31 ===
batch_distribution = pd.crosstab(
    sample_df["batch_number"],
    sample_df["sampling_group"],
)

batch_distribution["total"] = (
    batch_distribution.sum(axis=1)
)

display(batch_distribution)

if not (
    batch_distribution["total"]
    .eq(BATCH_SIZE)
    .all()
):
    raise ValueError(
        "Có batch không đủ đúng 50 bài."
    )

# === Notebook cell 32 ===
ANNOTATION_COLUMNS = [
    "promotion_label",
    "target_type",
    "target_entity_text",
    "target_company_id",
    "target_ticker",
    "sentiment",
    "esg_relevance",
    "environmental",
    "social",
    "governance",
    "esg_impact",
    "esg_information_type",
    "evidence_text",
    "annotator_note",
    "annotator_id",
    "annotation_status",
]

for column in ANNOTATION_COLUMNS:
    sample_df[column] = ""

sample_df["annotation_status"] = (
    "not_started"
)

# === Notebook cell 33 ===
DROPDOWN_VALUES = {
    "promotion_label": [
        "independent_editorial",
        "company_sourced",
        "promotional",
        "uncertain",
    ],
    "target_type": [
        "company",
        "industry",
        "market",
        "regulator",
        "project",
        "multiple",
        "other",
        "uncertain",
    ],
    "sentiment": [
        "positive",
        "negative",
        "mixed",
        "neutral",
    ],
    "esg_relevance": [
        "relevant",
        "not_relevant",
        "uncertain",
    ],
    "environmental": [
        "0",
        "1",
    ],
    "social": [
        "0",
        "1",
    ],
    "governance": [
        "0",
        "1",
    ],
    "esg_impact": [
        "opportunity",
        "risk",
        "mixed",
        "not_assessable",
        "not_applicable",
    ],
    "esg_information_type": [
        "aspiration",
        "commitment",
        "action",
        "outcome",
        "controversy",
        "mixed",
        "not_applicable",
    ],
    "annotation_status": [
        "not_started",
        "completed",
        "needs_review",
        "adjudicated",
    ],
}

OUTPUT_COLUMNS = [
    "article_id",
    "source",
    "published_at",
    "category",
    "title",
    "description",
    "content",
    "article_url",
    "promotion_label",
    "target_type",
    "target_entity_text",
    "target_company_id",
    "target_ticker",
    "sentiment",
    "esg_relevance",
    "environmental",
    "social",
    "governance",
    "esg_impact",
    "esg_information_type",
    "evidence_text",
    "annotator_note",
    "annotator_id",
    "annotation_status",
]

# === Notebook cell 34 ===
def excel_safe_text(
    value,
    max_length=MAX_CONTENT_LENGTH_EXCEL,
):
    if pd.isna(value):
        return ""

    value = str(value)

    if len(value) > max_length:
        return (
            value[:max_length]
            + "\n[CONTENT TRUNCATED]"
        )

    return value


def create_annotation_workbook(
    batch_df,
    output_path,
):
    wb = Workbook()

    # ========================================================
    # SHEET 1: INSTRUCTIONS
    # ========================================================

    ws_instruction = wb.active
    ws_instruction.title = "Instructions"

    instructions = [
        ["HƯỚNG DẪN GÁN NHÃN ESG NEWS"],
        [""],
        [
            "1",
            "Đọc title, description và content.",
        ],
        [
            "2",
            "Xác định đối tượng trung tâm của bài.",
        ],
        [
            "3",
            "Gán sentiment đối với đối tượng trung tâm.",
        ],
        [
            "4",
            "Xác định bài có thực sự liên quan ESG hay không.",
        ],
        [
            "5",
            "Nếu relevant, gán độc lập E, S và G.",
        ],
        [
            "6",
            "Nếu not_relevant, đặt E=S=G=0.",
        ],
        [
            "7",
            "Nếu not_relevant, đặt esg_impact và "
            "esg_information_type là not_applicable.",
        ],
        [
            "8",
            "Mixed chỉ dùng khi hai chiều hoặc hai loại "
            "thông tin đều có vai trò đáng kể.",
        ],
        [
            "9",
            "Uncertain phải giải thích trong annotator_note.",
        ],
        [
            "10",
            "Sao chép 1–3 câu quan trọng vào evidence_text.",
        ],
        [
            "11",
            "Không gán promotional tự động chỉ vì bài "
            "có nội dung tích cực.",
        ],
        [
            "12",
            "Chỉ chọn completed khi đã hoàn thành "
            "toàn bộ trường bắt buộc.",
        ],
    ]

    for row in instructions:
        ws_instruction.append(row)

    ws_instruction["A1"].font = Font(
        bold=True,
        size=16,
        color="FFFFFF",
    )

    ws_instruction["A1"].fill = PatternFill(
        "solid",
        fgColor="1F4E78",
    )

    ws_instruction.column_dimensions[
        "A"
    ].width = 12

    ws_instruction.column_dimensions[
        "B"
    ].width = 110

    # ========================================================
    # SHEET 2: ANNOTATION
    # ========================================================

    ws = wb.create_sheet(
        "Annotation"
    )

    for column_index, column_name in enumerate(
        OUTPUT_COLUMNS,
        start=1,
    ):
        cell = ws.cell(
            row=1,
            column=column_index,
            value=column_name,
        )

        cell.font = Font(
            bold=True,
            color="FFFFFF",
        )

        cell.fill = PatternFill(
            "solid",
            fgColor="2F75B5",
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    for row_index, (_, row) in enumerate(
        batch_df.iterrows(),
        start=2,
    ):
        for column_index, column_name in enumerate(
            OUTPUT_COLUMNS,
            start=1,
        ):
            value = row.get(
                column_name,
                "",
            )

            if (
                column_name == "published_at"
                and pd.notna(value)
            ):
                try:
                    value = pd.Timestamp(
                        value
                    ).strftime(
                        "%Y-%m-%d %H:%M:%S"
                    )
                except Exception:
                    value = str(value)

            ws.cell(
                row=row_index,
                column=column_index,
                value=excel_safe_text(value),
            )

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 35

    widths = {
        "article_id": 20,
        "source": 15,
        "published_at": 21,
        "category": 18,
        "title": 48,
        "description": 58,
        "content": 100,
        "article_url": 40,
        "promotion_label": 24,
        "target_type": 18,
        "target_entity_text": 30,
        "target_company_id": 20,
        "target_ticker": 15,
        "sentiment": 16,
        "esg_relevance": 18,
        "environmental": 16,
        "social": 12,
        "governance": 16,
        "esg_impact": 20,
        "esg_information_type": 24,
        "evidence_text": 65,
        "annotator_note": 45,
        "annotator_id": 16,
        "annotation_status": 20,
    }

    for index, column_name in enumerate(
        OUTPUT_COLUMNS,
        start=1,
    ):
        ws.column_dimensions[
            get_column_letter(index)
        ].width = widths.get(
            column_name,
            18,
        )

    thin_border = Border(
        left=Side(
            style="thin",
            color="D9E2F3",
        ),
        right=Side(
            style="thin",
            color="D9E2F3",
        ),
        top=Side(
            style="thin",
            color="D9E2F3",
        ),
        bottom=Side(
            style="thin",
            color="D9E2F3",
        ),
    )

    for row_cells in ws.iter_rows(
        min_row=2,
        max_row=ws.max_row,
    ):
        for cell in row_cells:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

            cell.border = thin_border

    # ========================================================
    # SHEET 3: DROPDOWN VALUES
    # ========================================================

    ws_dropdown = wb.create_sheet(
        "Dropdown_Values"
    )

    dropdown_ranges = {}

    for column_index, (
        field,
        values,
    ) in enumerate(
        DROPDOWN_VALUES.items(),
        start=1,
    ):
        ws_dropdown.cell(
            row=1,
            column=column_index,
            value=field,
        )

        for row_index, value in enumerate(
            values,
            start=2,
        ):
            ws_dropdown.cell(
                row=row_index,
                column=column_index,
                value=value,
            )

        column_letter = get_column_letter(
            column_index
        )

        dropdown_ranges[field] = (
            f"'Dropdown_Values'!"
            f"${column_letter}$2:"
            f"${column_letter}${len(values) + 1}"
        )

    output_positions = {
        name: position + 1
        for position, name
        in enumerate(OUTPUT_COLUMNS)
    }

    for field, formula_range in (
        dropdown_ranges.items()
    ):
        target_letter = get_column_letter(
            output_positions[field]
        )

        validation = DataValidation(
            type="list",
            formula1=formula_range,
            allow_blank=True,
        )

        validation.error = (
            "Hãy chọn giá trị trong danh sách."
        )
        validation.errorTitle = (
            "Giá trị không hợp lệ"
        )
        validation.prompt = (
            "Chọn một nhãn từ dropdown."
        )
        validation.promptTitle = field

        ws.add_data_validation(
            validation
        )

        validation.add(
            f"{target_letter}2:"
            f"{target_letter}{len(batch_df) + 1}"
        )

    ws_dropdown.sheet_state = "hidden"

    # ========================================================
    # SHEET 4: SAMPLING AUDIT
    # ========================================================

    ws_audit = wb.create_sheet(
        "Sampling_Audit"
    )

    audit_columns = [
        "article_id",
        "sampling_group",
        "esg_proxy",
        "promotion_proxy",
        "title_esg_proxy",
        "article_length",
        "year_group",
        "length_group",
        "input_file",
    ]

    for column_index, name in enumerate(
        audit_columns,
        start=1,
    ):
        audit_cell = ws_audit.cell(
            row=1,
            column=column_index,
            value=name,
        )

        audit_cell.font = Font(
            bold=True
        )

    for row_index, (_, row) in enumerate(
        batch_df.iterrows(),
        start=2,
    ):
        for column_index, name in enumerate(
            audit_columns,
            start=1,
        ):
            ws_audit.cell(
                row=row_index,
                column=column_index,
                value=excel_safe_text(
                    row.get(name, "")
                ),
            )

    # Không để annotator nhìn thấy proxy
    ws_audit.sheet_state = "hidden"

    wb.save(output_path)

# === Notebook cell 35 ===
old_batch_files = list(
    BATCH_DIR.glob(
        "annotation_batch_*.xlsx"
    )
)

print(
    "Số batch cũ sẽ thay thế:",
    len(old_batch_files),
)

for old_file in old_batch_files:
    old_file.unlink()

print("Đã dọn các batch cũ.")

# === Notebook cell 36 ===
batch_files = []

for batch_number in range(
    1,
    N_BATCHES + 1,
):
    batch_df = sample_df[
        sample_df["batch_number"].eq(
            batch_number
        )
    ].copy()

    if len(batch_df) != BATCH_SIZE:
        raise ValueError(
            f"Batch {batch_number} có "
            f"{len(batch_df)} bài, "
            f"không phải {BATCH_SIZE}."
        )

    output_path = (
        BATCH_DIR
        / (
            f"annotation_batch_"
            f"{batch_number:02d}.xlsx"
        )
    )

    create_annotation_workbook(
        batch_df=batch_df,
        output_path=output_path,
    )

    batch_files.append(
        output_path
    )

    print(
        f"Đã tạo {output_path.name}: "
        f"{len(batch_df)} bài"
    )

# === Notebook cell 37 ===
# Bổ sung nguồn dữ liệu cho các bài đã crawl
if "source" not in sample_df.columns:
    sample_df["source"] = "Vietstock"
else:
    sample_df["source"] = (
        sample_df["source"]
        .fillna("")
        .astype(str)
        .str.strip()
        .replace("", "Vietstock")
    )

print(sample_df["source"].value_counts(dropna=False))

# === Notebook cell 38 ===
master_parquet_path = (
    SAMPLE_DIR
    / "annotation_master_2000.parquet"
)

master_excel_path = (
    SAMPLE_DIR
    / "annotation_master_2000.xlsx"
)

sample_df.to_parquet(
    master_parquet_path,
    index=False,
)

master_export_columns = (
    OUTPUT_COLUMNS
    + [
        "batch_number",
        "sampling_group",
        "esg_proxy",
        "promotion_proxy",
        "article_length",
        "input_file",
    ]
)

sample_df[
    master_export_columns
].to_excel(
    master_excel_path,
    index=False,
)

print("Parquet:", master_parquet_path)
print("Excel  :", master_excel_path)

# === Notebook cell 39 ===
